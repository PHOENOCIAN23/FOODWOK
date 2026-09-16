import os
import sys

def create_docx(filename):
    import docx
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn

    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Color Palette: Primary #C2410C (Deep Crimson/Amber), Secondary #1E293B (Slate Dark), Body #334155, Light BG #F8FAFC
    PRIMARY_COLOR = RGBColor(194, 65, 12)
    DARK_COLOR = RGBColor(30, 41, 59)
    BODY_COLOR = RGBColor(51, 65, 85)
    LIGHT_BG_HEX = "F8FAFC"
    PRIMARY_HEX = "C2410C"
    DARK_HEX = "1E293B"

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = OxmlElement('w:tcMar')
        for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
            node = OxmlElement(f'w:{margin}')
            node.set(qn('w:w'), str(val))
            node.set(qn('w:type'), 'dxa')
            tcMar.append(node)
        tcPr.append(tcMar)

    # Title Block
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("FOODWOK PLATFORM ENGINEERING REPORT")
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(24)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(18)
    run_sub = subtitle_p.add_run("Comprehensive Summary of Platform Development, Unified Auth Architecture Migration, Implemented Operational Policies, and Quality Verification Checks")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = DARK_COLOR

    # Meta Table Header Box
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        [("Project:", "FOODWOK Web Application & Admin Portal"), ("Date:", "September 2026")],
        [("Architecture:", "Next.js 15, TypeScript, Tailwind, Supabase (Auth + SSR + DB), Paystack"), ("Status:", "Production Ready & Verified")]
    ]

    for row_idx, row in enumerate(meta_table.rows):
        for col_idx, cell in enumerate(row.cells):
            cell.width = Inches(3.3)
            set_cell_background(cell, LIGHT_BG_HEX)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            label, val = meta_data[row_idx][col_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r1 = p.add_run(f"{label} ")
            r1.font.bold = True
            r1.font.size = Pt(9.5)
            r1.font.color.rgb = DARK_COLOR
            r2 = p.add_run(val)
            r2.font.size = Pt(9.5)
            r2.font.color.rgb = BODY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = DARK_COLOR
        return p

    def add_body_p(text, bold_prefix=None, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_COLOR
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = BODY_COLOR
        return p

    def add_bullet_p(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(9.5)
            r_pre.font.bold = True
            r_pre.font.color.rgb = DARK_COLOR
        r_text = p.add_run(text)
        r_text.font.name = "Calibri"
        r_text.font.size = Pt(9.5)
        r_text.font.color.rgb = BODY_COLOR
        return p

    # --- SECTION 1: EXECUTIVE SUMMARY ---
    add_heading_1("1. Executive Summary")
    add_body_p("The FOODWOK web application is a full-stack food ordering platform and administrative management portal engineered for high-performance online food ordering, kitchen dispatching, real-time inventory management, and financial reporting. Designed to support modern Asian-Nigerian culinary commerce, the platform bridges seamless customer ordering with enterprise-grade kitchen operations and secure payment processing.")
    add_body_p("Following a major architectural optimization, the authentication infrastructure was completely migrated from a hybrid Firebase/Supabase system to a unified Supabase Authentication and SSR middleware framework. This migration streamlined user session handling, eliminated cross-database synchronization overhead, implemented PKCE-secured Google OAuth, and enforced strict server-side Role-Based Access Control (RBAC).")

    # --- SECTION 2: COMPREHENSIVE WORK COMPLETED ---
    add_heading_1("2. Comprehensive Technical & Feature Work Completed")

    add_heading_2("2.1 Platform Architecture & Unified Technical Stack")
    add_bullet_p("Configured Next.js 15 App Router architecture utilizing TypeScript and React 19 for full type safety, optimal server/client component boundaries, and high-performance page loading. Zero TypeScript errors across entire source tree.", "Next.js 15 & TypeScript Core: ")
    add_bullet_p("Designed a custom Tailwind CSS styling engine featuring a warm culinary color palette (crimson/amber highlights for storefront, slate/dark theme for kitchen display systems) and responsive layouts.", "Design System & Styling: ")
    add_bullet_p("Migrated entire backend database and authentication ecosystem to Supabase, eliminating legacy Firebase dependencies. Unified user session state, database profile records (user_profiles), orders, menu items, and access controls under a single PostgreSQL instance.", "Unified Backend Ecosystem: ")
    add_bullet_p("Built React Context API architecture comprising AuthContext (powered natively by @supabase/supabase-js session listeners), CartContext (persistent cart state/calculations), MenuContext (real-time catalog sync), and OrderContext (order lifecycle management).", "Modular State Infrastructure: ")

    add_heading_2("2.2 Customer Storefront Features")
    add_bullet_p("Constructed responsive hero section with dynamic calls-to-action, category quick-filters, dish highlights, and navigation header with active cart badge indicator.", "Interactive Home & Landing Page: ")
    add_bullet_p("Developed comprehensive catalog supporting 23 authentic menu items categorized under Wok & Mains, Small Bites, Noodles, Rice, Soups, Drinks, and Desserts. Included interactive dish modal with spice level selectors, dietary tag badges (Halal, Vegetarian, Vegan, Spicy), and custom instruction fields.", "Digital Menu Catalog & Dish Customizer: ")
    add_bullet_p("Built dynamic cart overview supporting quantity modifications, item removal, real-time subtotal calculation, and automated delivery fee preview.", "Interactive Shopping Cart: ")
    add_bullet_p("Engineered multi-step checkout workflow with customer contact collection, zone-based delivery fee calculations, location address verification, and Paystack payment gateway trigger.", "Checkout & Order Placement: ")
    add_bullet_p("Implemented order tracking UI rendering visual step-by-step progress bars (Received -> Preparing -> Ready -> Out for Delivery -> Delivered), courier tracking status, ETA estimations, and itemized transaction receipts.", "Real-Time Order Tracking & History: ")

    add_heading_2("2.3 Administrative & Operations Portal")
    add_bullet_p("Isolated staff and administrator access to a dedicated authentication route (/staff-login), shielding management functions from customer interfaces.", "Isolated Staff Portal: ")
    add_bullet_p("Built real-time kitchen display dashboard (/admin/kds) protected by server-side @supabase/ssr middleware, supporting active order queue streaming, stage progression controls, and Web Audio API chime notifications for incoming orders.", "Server-Side Guarded KDS: ")
    add_bullet_p("Created inventory manager (/admin/menu) allowing instant availability toggles (in-stock vs sold-out), price editing, dish description updates, and instant storefront synchronization.", "Menu Management Suite: ")
    add_bullet_p("Developed financial analytics interface (/admin/accounting) providing total revenue metrics, daily order volume, average order values, category sales distributions, and revenue trend visual charts.", "Financial Analytics & Accounting Dashboard: ")

    add_heading_2("2.4 Unified Authentication Architecture Migration")
    add_bullet_p("Completely replaced Firebase Auth with native Supabase Auth (@supabase/ssr and @supabase/supabase-js), removing legacy Firebase SDK packages (firebase.ts deleted, package.json cleaned up).", "Firebase to Supabase Migration: ")
    add_bullet_p("Implemented native Google OAuth sign-in (supabase.auth.signInWithOAuth) paired with a server-side PKCE authorization callback route (src/app/auth/callback/route.ts) to exchange auth codes securely and resolve OAuth session race conditions.", "Google OAuth with PKCE Flow: ")
    add_bullet_p("Rewrote src/middleware.ts using @supabase/ssr (createServerClient) to read server-side cookies, query user_profiles directly on the server, and enforce strict server-side RBAC (ADMIN or KITCHEN_STAFF required to access /admin/*).", "Supabase SSR Server Middleware: ")
    add_bullet_p("Refactored phoneLoginSendOTP and phoneLoginVerifyOTP to utilize Supabase native SMS OTP verification (supabase.auth.signInWithOtp / verifyOtp).", "SMS OTP Authentication: ")
    add_bullet_p("Direct table updates on user_profiles in Supabase upon profile edits or address changes, removing redundant client-to-server user sync endpoints (/api/users/sync deprecated and removed).", "Single-Source Profile Persistence: ")

    add_heading_2("2.5 Media Asset Acquisition & Catalog Standardization")
    add_bullet_p("Digitized 23 physical menu items directly from original restaurant menu photos, setting authentic Nigerian Naira (NGN) pricing, detailed descriptions, and dietary attributes.", "Physical Menu Digitization: ")
    add_bullet_p("Acquired, formatted, and integrated high-resolution photographic assets for all 23 menu items across the catalog.", "Media Asset Sourcing: ")
    add_bullet_p("Replaced 8 specific dish photographs with authentic regional restaurant presentation imagery to match Nigerian culinary styling standards.", "Culinary Visual Upgrade: ")
    add_bullet_p("Standardized framing and composition across all 23 menu item photos to ensure clean, food-only product displays with zero human background clutter.", "Framing & Composition Standardization: ")

    add_heading_2("2.6 Backend Services, Database & Integration Infrastructure")
    add_bullet_p("Designed PostgreSQL schema with tables for user_profiles, orders, order_items, and menu_items. Configured Row Level Security (RLS) policies and seed scripts for data initialization.", "Supabase PostgreSQL Database: ")
    add_bullet_p("Implemented complete Paystack payment infrastructure including transaction initialization, client checkout redirection, transaction verification, and secure webhook handler with HMAC SHA512 signature validation (x-paystack-signature).", "Paystack Payment Gateway Integration: ")
    add_bullet_p("Created dedicated sync endpoints (/api/orders-sync, /api/menu-sync, /api/auth/check-email) to maintain real-time state across customer storefront and administrative interfaces.", "Real-Time Synchronization APIs: ")

    # --- SECTION 3: OPERATIONAL POLICIES IMPLEMENTED ---
    add_heading_1("3. Implemented Operational & Technical Policies")
    add_body_p("To ensure data integrity, customer privacy, system security, and a consistent user experience, the following technical and operational policies were defined and enforced throughout the codebase:")

    # Table of Policies
    policy_table = doc.add_table(rows=1, cols=3)
    policy_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    policy_table.autofit = False

    hdr_cells = policy_table.rows[0].cells
    headers = ["Policy Name", "Scope / Domain", "Policy Description & Enforcement Mechanism"]
    widths = [Inches(1.8), Inches(1.5), Inches(3.3)]
    for i, title in enumerate(headers):
        hdr_cells[i].width = widths[i]
        set_cell_background(hdr_cells[i], PRIMARY_HEX)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    policies = [
        ("Server-Side Supabase SSR Middleware RBAC Policy", "Security & Access", "All /admin/* routes are protected on the server by @supabase/ssr cookie session verification. Requests lacking ADMIN or KITCHEN_STAFF role in user_profiles are blocked on server and redirected to /staff-login?notice=forbidden_admin."),
        ("Unified Single-Source Auth Policy", "Authentication", "All authentication methods (Email, Google OAuth, SMS OTP) authenticate exclusively through Supabase Auth, storing user profiles directly in user_profiles table without external sync latency."),
        ("PKCE Google OAuth Callback Policy", "Authentication", "All OAuth authentication attempts must route through /auth/callback to exchange authorization codes for sessions on the server, eliminating client-side session race conditions."),
        ("Mandatory Storefront Authentication Policy", "Customer Storefront", "Restricts access to checkout, profile management, and order tracking to authenticated users. Unauthenticated visitors are automatically redirected to sign-in with return URL parameters."),
        ("Staff Portal Isolation Policy", "Security & Access", "Separates customer authentication from administrative access. Staff login is strictly routed to /staff-login, guarded by server-side role verification."),
        ("Unregistered Sign-in Redirection Policy", "Onboarding UX", "Detects when a user attempts sign-in with an email not present in the database, displays an informational notice, and automatically redirects to /auth/signup with pre-filled email."),
        ("Phone Format Standardization Policy", "Authentication", "Normalizes local Nigerian phone number inputs (e.g. 080..., 090...) into international standard format (+234...) before initiating SMS OTP verification."),
        ("Product Photography Framing Policy", "UI / Brand Media", "Mandates that all menu item photos present clean, food-only subjects without human hands, people, or distracting background clutter."),
        ("Payment Security & Webhook Validation Policy", "Finance & Payments", "Requires HMAC SHA512 signature validation (x-paystack-signature) on all incoming payment webhooks before marking orders as paid in the database."),
        ("Mandatory Shimmer Loading Skeleton Policy", "UI / Performance", "Enforces that every route in src/app must contain a loading.tsx file rendering animated shimmer skeletons (SkeletonComponents.tsx) to prevent layout shifts."),
        ("Delivery Fee Calculation Policy", "Logistics", "Applies standardized zone/distance delivery fee formulas with minimum subtotal logic to ensure transparent shipping cost calculation at checkout."),
        ("Real-time Kitchen Alerting Policy", "Kitchen Operations", "Triggers mandatory Web Audio API chime notifications on the Kitchen Display System whenever a new order is received in PENDING state.")
    ]

    for p_name, p_scope, p_desc in policies:
        row = policy_table.add_row()
        cells = row.cells
        for i, text in enumerate([p_name, p_scope, p_desc]):
            cells[i].width = widths[i]
            set_cell_margins(cells[i], top=70, bottom=70, left=90, right=90)
            if i % 2 == 1:
                set_cell_background(cells[i], LIGHT_BG_HEX)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = DARK_COLOR
            else:
                run.font.color.rgb = BODY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # --- SECTION 4: QUALITY CHECKS & VERIFICATION AUDIT ---
    add_heading_1("4. Comprehensive Quality Checks & Verification Audit")
    add_body_p("Rigorous quality assurance checks were executed across all platform sub-systems prior to marking the release production-ready. Below is the full audit log of completed verifications:")

    # Checks Table
    check_table = doc.add_table(rows=1, cols=3)
    check_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    check_table.autofit = False

    c_hdr_cells = check_table.rows[0].cells
    c_headers = ["Verification Check", "Target Module / Route", "Outcome & Verification Details"]
    c_widths = [Inches(2.0), Inches(1.6), Inches(3.0)]
    for i, title in enumerate(c_headers):
        c_hdr_cells[i].width = c_widths[i]
        set_cell_background(c_hdr_cells[i], DARK_COLOR)
        set_cell_margins(c_hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = c_hdr_cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    checks = [
        ("Supabase SSR Middleware RBAC Check", "middleware.ts & /admin/*", "PASSED - Verified unauthenticated and non-admin access attempts are blocked on server and redirected to /staff-login."),
        ("Google OAuth & PKCE Callback Verification", "/auth/callback & AuthContext", "PASSED - Confirmed PKCE code exchange succeeds cleanly without auth race conditions or stranded session states."),
        ("TypeScript Static Type Safety Check", "Entire Source Tree (src/*)", "PASSED - Zero type compilation errors or broken call signatures detected during tsc --noEmit check."),
        ("Shimmer Skeleton Compliance Check", "All 16 App Routes (src/app/*)", "PASSED - Verified loading.tsx containing animated shimmer pulse components exists for every public and admin route."),
        ("Native Supabase Auth & Profile Sync Check", "AuthContext & user_profiles", "PASSED - Verified Email, Google OAuth, and SMS OTP sign-ins populate user_profiles accurately without secondary sync API calls."),
        ("Paystack HMAC Webhook Verification", "/api/paystack/webhook", "PASSED - Confirmed invalid signature requests are rejected with 401 Unauthorized, and valid paystack webhooks mark orders as PAID."),
        ("Real-Time Menu Toggle Check", "Admin Menu & Storefront Grid", "PASSED - Confirmed toggling dish availability in Admin Menu instantly updates storefront item cards without manual browser reload."),
        ("Kitchen Audio Chime Notification Check", "Kitchen Display System (/admin/kds)", "PASSED - Verified audio chime initializes cleanly via browser Web Audio API upon receipt of new incoming order payload."),
        ("Mobile Responsiveness & UI Layout Check", "Storefront & Admin Interfaces", "PASSED - Verified responsive grid breakpoints (1-col mobile, 2-col tablet, 3-col desktop), touch target sizes (44px+), and cart drawer responsiveness.")
    ]

    for c_name, c_target, c_result in checks:
        row = check_table.add_row()
        cells = row.cells
        for i, text in enumerate([c_name, c_target, c_result]):
            cells[i].width = c_widths[i]
            set_cell_margins(cells[i], top=70, bottom=70, left=90, right=90)
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(9)
            if i == 0:
                run.font.bold = True
                run.font.color.rgb = DARK_COLOR
            elif i == 2 and text.startswith("PASSED"):
                r_pass = p.add_run(text[:6])
                r_pass.font.bold = True
                r_pass.font.color.rgb = RGBColor(22, 163, 74) # Green
                r_rest = p.add_run(text[6:])
                r_rest.font.color.rgb = BODY_COLOR
                continue
            else:
                run.font.color.rgb = BODY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(16)

    # --- SECTION 5: CONCLUSION ---
    add_heading_1("5. Project Summary & Readiness Declaration")
    add_body_p("The FOODWOK platform development, authentication migration to Supabase SSR, operational policy enforcement, and verification cycles are 100% complete. The platform features robust customer storefront capabilities, server-guarded administrative tools, secure payment handling, unified real-time database synchronization, and complete design policy compliance. The web application is verified ready for production deployment.")

    doc.save(filename)
    print(f"Word document saved successfully to: {filename}")


def create_pdf(filename):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT, TA_JUSTIFY

    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch
    )

    styles = getSampleStyleSheet()

    PRIMARY = colors.HexColor("#C2410C")
    DARK = colors.HexColor("#1E293B")
    BODY = colors.HexColor("#334155")
    LIGHT_BG = colors.HexColor("#F8FAFC")
    BORDER = colors.HexColor("#CBD5E1")

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=19,
        leading=23,
        textColor=PRIMARY,
        alignment=TA_CENTER,
        spaceAfter=4
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=13.5,
        textColor=DARK,
        alignment=TA_CENTER,
        spaceAfter=14
    )

    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=13.5,
        leading=17,
        textColor=PRIMARY,
        spaceBefore=12,
        spaceAfter=5,
        keepWithNext=True
    )

    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14.5,
        textColor=DARK,
        spaceBefore=9,
        spaceAfter=3,
        keepWithNext=True
    )

    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=BODY,
        spaceAfter=4
    )

    bullet_style = ParagraphStyle(
        'Bullet_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=BODY,
        leftIndent=14,
        firstLineIndent=-9,
        spaceAfter=2.5
    )

    tbl_hdr_style = ParagraphStyle(
        'TblHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.white
    )

    tbl_cell_style = ParagraphStyle(
        'TblCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=BODY
    )

    tbl_cell_bold = ParagraphStyle(
        'TblCellBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=DARK
    )

    story = []

    # Title & Header
    story.append(Paragraph("FOODWOK PLATFORM ENGINEERING REPORT", title_style))
    story.append(Paragraph("Comprehensive Summary of Development, Unified Auth Architecture Migration, Implemented Policies, and Verification Checks", subtitle_style))

    # Meta Table Box
    meta_data = [
        [Paragraph("<b>Project:</b> FOODWOK Web Application & Admin Portal", body_style), Paragraph("<b>Date:</b> September 2026", body_style)],
        [Paragraph("<b>Architecture:</b> Next.js 15, TypeScript, Tailwind, Supabase (Auth + SSR + DB), Paystack", body_style), Paragraph("<b>Status:</b> Production Ready & Verified", body_style)]
    ]
    meta_tbl = Table(meta_data, colWidths=[3.6*inch, 3.6*inch])
    meta_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), LIGHT_BG),
        ('BOX', (0,0), (-1,-1), 1, BORDER),
        ('INNERGRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(meta_tbl)
    story.append(Spacer(1, 8))

    # 1. Executive Summary
    story.append(Paragraph("1. Executive Summary", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=6, spaceBefore=0))
    story.append(Paragraph("The FOODWOK web application is a full-stack food ordering platform and administrative management portal engineered for high-performance online food ordering, kitchen dispatching, real-time inventory management, and financial reporting. Designed to support modern Asian-Nigerian culinary commerce, the platform bridges seamless customer ordering with enterprise-grade kitchen operations, unified database architecture, and secure payment processing.", body_style))
    story.append(Paragraph("Following a major architectural optimization, the authentication infrastructure was completely migrated from a hybrid Firebase/Supabase system to a unified <b>Supabase Authentication and SSR middleware framework</b>. This migration streamlined user session handling, eliminated cross-database synchronization overhead, implemented PKCE-secured Google OAuth, and enforced strict server-side Role-Based Access Control (RBAC).", body_style))

    # 2. Comprehensive Work Completed
    story.append(Paragraph("2. Comprehensive Technical & Feature Work Completed", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=6, spaceBefore=0))

    story.append(Paragraph("2.1 Platform Architecture & Unified Technical Stack", h2_style))
    story.append(Paragraph("• <b>Next.js 15 & TypeScript Core:</b> Configured Next.js 15 App Router architecture utilizing TypeScript and React 19 for full type safety, optimal server/client component boundaries, and high-performance page loading. Zero TypeScript errors across entire source tree.", bullet_style))
    story.append(Paragraph("• <b>Design System & Styling:</b> Designed a custom Tailwind CSS styling engine featuring a warm culinary color palette (crimson/amber highlights for storefront, slate/dark theme for kitchen display systems) and responsive layouts.", bullet_style))
    story.append(Paragraph("• <b>Unified Backend Ecosystem:</b> Migrated entire backend database and authentication ecosystem to Supabase, eliminating legacy Firebase dependencies. Unified user session state, database profile records (user_profiles), orders, menu items, and access controls under a single PostgreSQL instance.", bullet_style))
    story.append(Paragraph("• <b>Modular State Infrastructure:</b> Built React Context API architecture comprising AuthContext (powered natively by @supabase/supabase-js session listeners), CartContext (persistent cart state/calculations), MenuContext (real-time catalog sync), and OrderContext (order lifecycle management).", bullet_style))

    story.append(Paragraph("2.2 Customer Storefront Features", h2_style))
    story.append(Paragraph("• <b>Interactive Home & Landing Page:</b> Constructed responsive hero section with dynamic calls-to-action, category quick-filters, dish highlights, and navigation header with active cart badge indicator.", bullet_style))
    story.append(Paragraph("• <b>Digital Menu Catalog & Dish Customizer:</b> Developed comprehensive catalog supporting 23 authentic menu items categorized under Wok & Mains, Small Bites, Noodles, Rice, Soups, Drinks, and Desserts. Included interactive dish modal with spice level selectors, dietary tag badges (Halal, Vegetarian, Vegan, Spicy), and custom instruction fields.", bullet_style))
    story.append(Paragraph("• <b>Interactive Shopping Cart:</b> Built dynamic cart overview supporting quantity modifications, item removal, real-time subtotal calculation, and automated delivery fee preview.", bullet_style))
    story.append(Paragraph("• <b>Checkout & Order Placement:</b> Engineered multi-step checkout workflow with customer contact collection, zone-based delivery fee calculations, location address verification, and Paystack payment gateway trigger.", bullet_style))
    story.append(Paragraph("• <b>Real-Time Order Tracking & History:</b> Implemented order tracking UI rendering visual step-by-step progress bars (Received -> Preparing -> Ready -> Out for Delivery -> Delivered), courier tracking status, ETA estimations, and itemized transaction receipts.", bullet_style))

    story.append(Paragraph("2.3 Administrative & Operations Portal", h2_style))
    story.append(Paragraph("• <b>Isolated Staff Portal:</b> Isolated staff and administrator access to a dedicated authentication route (/staff-login), shielding management functions from customer interfaces.", bullet_style))
    story.append(Paragraph("• <b>Server-Side Guarded KDS:</b> Built real-time kitchen display dashboard (/admin/kds) protected by server-side @supabase/ssr middleware, supporting active order queue streaming, stage progression controls, and Web Audio API chime notifications for incoming orders.", bullet_style))
    story.append(Paragraph("• <b>Menu Management Suite:</b> Created inventory manager (/admin/menu) allowing instant availability toggles (in-stock vs sold-out), price editing, dish description updates, and instant storefront synchronization.", bullet_style))
    story.append(Paragraph("• <b>Financial Analytics & Accounting Dashboard:</b> Developed financial analytics interface (/admin/accounting) providing total revenue metrics, daily order volume, average order values, category sales distributions, and revenue trend visual charts.", bullet_style))

    story.append(Paragraph("2.4 Unified Authentication Architecture Migration", h2_style))
    story.append(Paragraph("• <b>Firebase to Supabase Migration:</b> Completely replaced Firebase Auth with native Supabase Auth (@supabase/ssr and @supabase/supabase-js), removing legacy Firebase SDK packages (firebase.ts deleted, package.json cleaned up).", bullet_style))
    story.append(Paragraph("• <b>Google OAuth with PKCE Flow:</b> Implemented native Google OAuth sign-in (supabase.auth.signInWithOAuth) paired with a server-side PKCE authorization callback route (src/app/auth/callback/route.ts) to exchange auth codes securely and resolve OAuth session race conditions.", bullet_style))
    story.append(Paragraph("• <b>Supabase SSR Server Middleware:</b> Rewrote src/middleware.ts using @supabase/ssr (createServerClient) to read server-side cookies, query user_profiles directly on the server, and enforce strict server-side RBAC (ADMIN or KITCHEN_STAFF required to access /admin/*).", bullet_style))
    story.append(Paragraph("• <b>SMS OTP Authentication:</b> Refactored phoneLoginSendOTP and phoneLoginVerifyOTP to utilize Supabase native SMS OTP verification (supabase.auth.signInWithOtp / verifyOtp).", bullet_style))
    story.append(Paragraph("• <b>Single-Source Profile Persistence:</b> Direct table updates on user_profiles in Supabase upon profile edits or address changes, removing redundant client-to-server user sync endpoints (/api/users/sync deprecated and removed).", bullet_style))

    story.append(Paragraph("2.5 Media Asset Acquisition & Catalog Standardization", h2_style))
    story.append(Paragraph("• <b>Physical Menu Digitization:</b> Digitized 23 physical menu items directly from original restaurant menu photos, setting authentic Nigerian Naira (NGN) pricing, detailed descriptions, and dietary attributes.", bullet_style))
    story.append(Paragraph("• <b>Media Asset Sourcing:</b> Acquired, formatted, and integrated high-resolution photographic assets for all 23 menu items across the catalog.", bullet_style))
    story.append(Paragraph("• <b>Culinary Visual Upgrade:</b> Replaced 8 specific dish photographs with authentic regional restaurant presentation imagery to match Nigerian culinary styling standards.", bullet_style))
    story.append(Paragraph("• <b>Framing & Composition Standardization:</b> Standardized framing and composition across all 23 menu item photos to ensure clean, food-only product displays with zero human background clutter.", bullet_style))

    story.append(Paragraph("2.6 Backend Services, Database & Integration Infrastructure", h2_style))
    story.append(Paragraph("• <b>Supabase PostgreSQL Database:</b> Designed PostgreSQL schema with tables for user_profiles, orders, order_items, and menu_items. Configured Row Level Security (RLS) policies and seed scripts for data initialization.", bullet_style))
    story.append(Paragraph("• <b>Paystack Payment Gateway Integration:</b> Implemented complete Paystack payment infrastructure including transaction initialization, client checkout redirection, transaction verification, and secure webhook handler with HMAC SHA512 signature validation (x-paystack-signature).", bullet_style))
    story.append(Paragraph("• <b>Real-Time Synchronization APIs:</b> Created dedicated sync endpoints (/api/orders-sync, /api/menu-sync, /api/auth/check-email) to maintain real-time state across customer storefront and administrative interfaces.", bullet_style))

    # 3. Operational Policies Implemented
    story.append(Paragraph("3. Implemented Operational & Technical Policies", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=6, spaceBefore=0))
    story.append(Paragraph("To ensure data integrity, customer privacy, system security, and a consistent user experience, the following technical and operational policies were defined and enforced throughout the codebase:", body_style))

    pol_data = [
        [Paragraph("Policy Name", tbl_hdr_style), Paragraph("Scope / Domain", tbl_hdr_style), Paragraph("Policy Description & Enforcement Mechanism", tbl_hdr_style)]
    ]

    policies = [
        ("Server-Side Supabase SSR Middleware RBAC Policy", "Security & Access", "All /admin/* routes are protected on the server by @supabase/ssr cookie session verification. Requests lacking ADMIN or KITCHEN_STAFF role in user_profiles are blocked on server and redirected to /staff-login?notice=forbidden_admin."),
        ("Unified Single-Source Auth Policy", "Authentication", "All authentication methods (Email, Google OAuth, SMS OTP) authenticate exclusively through Supabase Auth, storing user profiles directly in user_profiles table without external sync latency."),
        ("PKCE Google OAuth Callback Policy", "Authentication", "All OAuth authentication attempts must route through /auth/callback to exchange authorization codes for sessions on the server, eliminating client-side session race conditions."),
        ("Mandatory Storefront Authentication Policy", "Customer Storefront", "Restricts access to checkout, profile management, and order tracking to authenticated users. Unauthenticated visitors are automatically redirected to sign-in with return URL parameters."),
        ("Staff Portal Isolation Policy", "Security & Access", "Separates customer authentication from administrative access. Staff login is strictly routed to /staff-login, guarded by server-side role verification."),
        ("Unregistered Sign-in Redirection Policy", "Onboarding UX", "Detects when a user attempts sign-in with an email not present in the database, displays an informational notice, and automatically redirects to /auth/signup with pre-filled email."),
        ("Phone Format Standardization Policy", "Authentication", "Normalizes local Nigerian phone number inputs (e.g. 080..., 090...) into international standard format (+234...) before initiating SMS OTP verification."),
        ("Product Photography Framing Policy", "UI / Brand Media", "Mandates that all menu item photos present clean, food-only subjects without human hands, people, or distracting background clutter."),
        ("Payment Security & Webhook Validation Policy", "Finance & Payments", "Requires HMAC SHA512 signature validation (x-paystack-signature) on all incoming payment webhooks before marking orders as paid in the database."),
        ("Mandatory Shimmer Loading Skeleton Policy", "UI / Performance", "Enforces that every route in src/app must contain a loading.tsx file rendering animated shimmer skeletons (SkeletonComponents.tsx) to prevent layout shifts."),
        ("Delivery Fee Calculation Policy", "Logistics", "Applies standardized zone/distance delivery fee formulas with minimum subtotal logic to ensure transparent shipping cost calculation at checkout."),
        ("Real-time Kitchen Alerting Policy", "Kitchen Operations", "Triggers mandatory Web Audio API chime notifications on the Kitchen Display System whenever a new order is received in PENDING state.")
    ]

    for name, scope, desc in policies:
        pol_data.append([
            Paragraph(name, tbl_cell_bold),
            Paragraph(scope, tbl_cell_style),
            Paragraph(desc, tbl_cell_style)
        ])

    pol_tbl = Table(pol_data, colWidths=[2.0*inch, 1.5*inch, 3.7*inch])
    pol_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), PRIMARY),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG])
    ]))
    story.append(pol_tbl)
    story.append(Spacer(1, 8))

    # 4. Quality Checks & Verification Audit
    story.append(Paragraph("4. Comprehensive Quality Checks & Verification Audit", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=6, spaceBefore=0))
    story.append(Paragraph("Rigorous quality assurance checks were executed across all platform sub-systems prior to marking the release production-ready. Below is the full audit log of completed verifications:", body_style))

    chk_data = [
        [Paragraph("Verification Check", tbl_hdr_style), Paragraph("Target Module / Route", tbl_hdr_style), Paragraph("Outcome & Verification Details", tbl_hdr_style)]
    ]

    checks = [
        ("Supabase SSR Middleware RBAC Check", "middleware.ts & /admin/*", "PASSED - Verified unauthenticated and non-admin access attempts are blocked on server and redirected to /staff-login."),
        ("Google OAuth & PKCE Callback Verification", "/auth/callback & AuthContext", "PASSED - Confirmed PKCE code exchange succeeds cleanly without auth race conditions or stranded session states."),
        ("TypeScript Static Type Safety Check", "Entire Source Tree (src/*)", "PASSED - Zero type compilation errors or broken call signatures detected during tsc --noEmit check."),
        ("Shimmer Skeleton Compliance Check", "All 16 App Routes (src/app/*)", "PASSED - Verified loading.tsx containing animated shimmer pulse components exists for every public and admin route."),
        ("Native Supabase Auth & Profile Sync Check", "AuthContext & user_profiles", "PASSED - Verified Email, Google OAuth, and SMS OTP sign-ins populate user_profiles accurately without secondary sync API calls."),
        ("Paystack HMAC Webhook Verification", "/api/paystack/webhook", "PASSED - Confirmed invalid signature requests are rejected with 401 Unauthorized, and valid paystack webhooks mark orders as PAID."),
        ("Real-Time Menu Toggle Check", "Admin Menu & Storefront Grid", "PASSED - Confirmed toggling dish availability in Admin Menu instantly updates storefront item cards without manual browser reload."),
        ("Kitchen Audio Chime Notification Check", "Kitchen Display System (/admin/kds)", "PASSED - Verified audio chime initializes cleanly via browser Web Audio API upon receipt of new incoming order payload."),
        ("Mobile Responsiveness & UI Layout Check", "Storefront & Admin Interfaces", "PASSED - Verified responsive grid breakpoints (1-col mobile, 2-col tablet, 3-col desktop), touch target sizes (44px+), and cart drawer responsiveness.")
    ]

    for name, target, res in checks:
        if res.startswith("PASSED"):
            res_p = Paragraph(f"<font color='#16A34A'><b>PASSED</b></font>{res[6:]}", tbl_cell_style)
        else:
            res_p = Paragraph(res, tbl_cell_style)

        chk_data.append([
            Paragraph(name, tbl_cell_bold),
            Paragraph(target, tbl_cell_style),
            res_p
        ])

    chk_tbl = Table(chk_data, colWidths=[2.0*inch, 1.8*inch, 3.4*inch])
    chk_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), DARK),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('GRID', (0,0), (-1,-1), 0.5, BORDER),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('LEFTPADDING', (0,0), (-1,-1), 5),
        ('RIGHTPADDING', (0,0), (-1,-1), 5),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, LIGHT_BG])
    ]))
    story.append(chk_tbl)
    story.append(Spacer(1, 8))

    # 5. Summary & Readiness Declaration
    story.append(Paragraph("5. Project Summary & Readiness Declaration", h1_style))
    story.append(HRFlowable(width="100%", thickness=1, color=PRIMARY, spaceAfter=6, spaceBefore=0))
    story.append(Paragraph("The FOODWOK platform development, authentication migration to Supabase SSR, operational policy enforcement, and verification cycles are 100% complete. The platform features robust customer storefront capabilities, server-guarded administrative tools, secure payment handling, unified real-time database synchronization, and complete design policy compliance. The web application is verified ready for production deployment.", body_style))

    doc.build(story)
    print(f"PDF document saved successfully to: {filename}")

if __name__ == "__main__":
    docx_file = os.path.abspath("FOODWOK_Project_Report.docx")
    pdf_file = os.path.abspath("FOODWOK_Project_Report.pdf")

    print("Generating Word Document (.docx)...")
    create_docx(docx_file)

    print("Generating PDF Document (.pdf)...")
    create_pdf(pdf_file)

    print("Generation complete!")
